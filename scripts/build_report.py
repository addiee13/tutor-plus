from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

from report_content import BACKLOG, CLIENT_REQUIREMENTS, PLAN, REQUIREMENTS, SPRINT1_CHANGES, USE_CASES


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output/doc"
OUT.mkdir(parents=True, exist_ok=True)
TOC_FILE = OUT / "toc.json"
FIGURES = ROOT / "docs/diagrams/rendered"
CLOUD = ROOT / "docs/evidence/screenshots/cloud"
DOC = Document()
FIGURE_NUMBER = 0
TABLE_NUMBER = 0
HEADINGS = []
TOC_PAGES = {}


def inline(paragraph, text):
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    for piece in re.split(r"(\*\*.*?\*\*|`[^`]+`)", text):
        if not piece:
            continue
        run = paragraph.add_run(piece[2:-2] if piece.startswith("**") else piece.strip("`"))
        run.bold = piece.startswith("**")
        if piece.startswith("`"):
            run.font.name = "Courier New"
            run.font.size = Pt(9.5)


def body(text, style=None):
    paragraph = DOC.add_paragraph(style=style)
    inline(paragraph, text)
    return paragraph


def heading(text, level=1, new_page=False):
    paragraph = DOC.add_heading(text, level=level)
    paragraph.paragraph_format.page_break_before = new_page
    if level <= 2 and text != "TABLE OF CONTENTS":
        HEADINGS.append([level, text])
    return paragraph


def orientation(landscape):
    section = DOC.add_section(WD_SECTION_START.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Inches(11 if landscape else 8.5)
    section.page_height = Inches(8.5 if landscape else 11)
    section.header.is_linked_to_previous = True
    section.footer.is_linked_to_previous = True
    section.top_margin = Inches(0.35 if landscape else 0.85)
    section.bottom_margin = Inches(0.35 if landscape else 0.8)
    section.header_distance = Inches(0.12 if landscape else 0.4)
    section.footer_distance = Inches(0.12 if landscape else 0.4)


def caption(text):
    paragraph = DOC.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inline(paragraph, text)
    return paragraph


def field(paragraph, instruction, value="1"):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    command = OxmlElement("w:instrText")
    command.set(qn("xml:space"), "preserve")
    command.text = instruction
    run._r.append(command)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.append(separate)
    run.add_text(value)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def configure():
    section = DOC.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(0.95)
    section.right_margin = Inches(0.95)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.4)
    for style_name in ["Normal", "Body Text", "List Bullet", "List Number"]:
        style = DOC.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(11)
        style.paragraph_format.line_spacing = 1.1
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.widow_control = True
    for level, size in [(1, 14), (2, 12), (3, 11.5), (4, 11)]:
        style = DOC.styles[f"Heading {level}"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    DOC.styles["Caption"].font.name = "Times New Roman"
    DOC.styles["Caption"].font.size = Pt(10)
    DOC.styles["Caption"].font.italic = False
    DOC.styles["Caption"].font.color.rgb = RGBColor(0, 0, 0)
    DOC.styles["Caption"].paragraph_format.space_after = Pt(8)
    header = section.header.paragraphs[0]
    header.paragraph_format.tab_stops.add_tab_stop(Inches(6.6), WD_TAB_ALIGNMENT.RIGHT)
    header.add_run("Group 11\tSoftware Engineering Project").font.size = Pt(10)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in [("val", "single"), ("sz", "4"), ("space", "1"), ("color", "000000")]:
        bottom.set(qn(f"w:{key}"), value)
    border.append(bottom)
    header._p.get_or_add_pPr().append(border)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(footer, " PAGE ")
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    DOC.settings.element.append(update)
    DOC.core_properties.title = "Tutor Plus: Smart Platform for Academic Progress, Sprint 2"
    DOC.core_properties.author = "Adam Dalwai; Shadmaan Akhand; Edu M. Okpok; Kien Duong; Nikhil Rodrigo"
    DOC.core_properties.subject = "Group 11 Software Engineering Fall 2026"


def table(headers, rows, widths=None, size=10):
    result = DOC.add_table(rows=1, cols=len(headers))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    result.style = "Table Grid"
    available_width = DOC.sections[-1].page_width.inches - 1.9
    widths = widths or [available_width / len(headers)] * len(headers)
    result.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for cell, value in zip(result.rows[0].cells, headers):
        cell.paragraphs[0].add_run(value).bold = True
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "EDEDED")
        cell._tc.get_or_add_tcPr().append(shading)
    for values in rows:
        for cell, value in zip(result.add_row().cells, values):
            cell.paragraphs[0].add_run(str(value))
    for row in result.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(size)
    DOC.add_paragraph().paragraph_format.space_after = Pt(2)


def figure(filename, description=None, maximum_height=7.3, directory=FIGURES, maximum_width=None):
    path = directory / filename
    if not path.exists():
        raise FileNotFoundError(f"Required figure missing: {path}")
    image = Image.open(path)
    available_width = maximum_width or DOC.sections[-1].page_width.inches - 1.9
    width = min(available_width, maximum_height * image.width / image.height)
    paragraph = DOC.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = bool(description)
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    if description:
        caption(description)


def labeled(label, text="", keep=False, left=False):
    """One 'Label: text' line, as in the Sprint 5 reference report."""
    paragraph = DOC.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.keep_with_next = keep
    if left:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.add_run(f"{label} ").bold = True
    paragraph.add_run(text)
    return paragraph


def numbered(items, keep=False, left=False):
    for number, item in enumerate(items, 1):
        paragraph = DOC.add_paragraph(f"{number}. {item}")
        paragraph.paragraph_format.keep_with_next = keep
        if left:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = Inches(0.45)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(1)


def bullets(items):
    for item in items:
        DOC.add_paragraph(item, style="List Bullet").paragraph_format.space_after = Pt(2)


def monospace(lines):
    for line in lines:
        paragraph = DOC.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = Inches(0.4)
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(9.5)
    DOC.add_paragraph().paragraph_format.space_after = Pt(2)


def database_rows():
    """Read the exact attributes and keys from the local catalog dictionary."""
    rows = {}
    for line in (ROOT / "docs/architecture/SPRINT2_DATABASE.md").read_text().splitlines():
        if not line.startswith(("| public.", "| private.")):
            continue
        cells = line.strip().strip("|").split(" | ")
        name = cells[0].strip().removeprefix("public.").removeprefix("private.")
        primary_keys, foreign_keys = set(), {}
        for part in cells[2].split("; "):
            primary = re.search(r"PRIMARY KEY \(([^)]+)\)", part)
            foreign = re.search(r"FOREIGN KEY \(([^)]+)\) REFERENCES ([\w.]+)\(([^)]+)\)", part)
            if primary:
                primary_keys.update(column.strip() for column in primary.group(1).split(","))
            if foreign:
                for column, target in zip(foreign.group(1).split(","), foreign.group(3).split(",")):
                    foreign_keys.setdefault(column.strip(), []).append(f"{foreign.group(2).removeprefix('public.')}.{target.strip()}")
        attributes = []
        for item in re.findall(r"`([^`]+)`", cells[1]):
            column, kind = item.split(" ", 1)
            label = ", ".join(
                (["PK"] if column in primary_keys else []) + [f"FK to {target}" for target in foreign_keys.get(column, [])]
            )
            attributes.append([column, kind.rstrip("?").replace("timestamp with time zone", "timestamptz"), label])
        rows[name] = attributes
    return rows


DATABASE_GROUPS = [
    ("UC-01 User Management", [
        ("profiles", "Private profile for each account. Its id matches the managed Supabase Auth user."),
        ("user_roles", "Student, Tutor, and Administrator capabilities for each account."),
        ("audit_logs", "Record of every administrator action and the reason given."),
        ("admin_bootstrap_guard", "Private table that lets the first Administrator be created only once."),
    ]),
    ("UC-02 Manage Student and Tutor Profiles", [
        ("subjects", "Catalog of subjects a tutor can teach."),
        ("tutor_profiles", "Tutor's own profile, rate, modes, and publication state."),
        ("tutor_subjects", "Links a tutor to the subjects they teach."),
        ("public_tutor_snapshots", "Public copy of a published tutor. It holds no private data."),
    ]),
    ("UC-04 Availability and Booking", [
        ("availability_rules", "Weekly times a tutor is available."),
        ("availability_exceptions", "Days off or extra times for a tutor."),
        ("bookings", "A reserved session. Exclusion constraints block overlapping tutor or student times."),
        ("booking_private_details", "Meeting link or place, visible only to the two people on the booking."),
        ("booking_events", "History of each booking status change."),
        ("booking_policy", "Private booking rules such as lead time, buffer, hold time, and platform fee."),
    ]),
    ("UC-06 Payments and Refunds", [
        ("payments", "Payment for one booking, stored in whole cents with the platform fee."),
        ("payment_events", "Verified provider events. Each event is processed once."),
        ("refunds", "Refund requests and their status."),
        ("payout_refs", "Reference to the tutor payout for a payment."),
    ]),
]


SKIP_SHOTS = {"01", "09", "14"}


def cover():
    DOC.add_paragraph().paragraph_format.space_after = Pt(90)
    DOC.add_paragraph().add_run("(Tutor Plus: Smart Platform for Academic Progress)").bold = True
    DOC.add_paragraph("Software Engineering Project").alignment = WD_ALIGN_PARAGRAPH.LEFT
    DOC.add_paragraph().paragraph_format.space_after = Pt(60)
    for text in [
        "Name (s): Adam Dalwai, Shadmaan Akhand, Edu M. Okpok, Kien Duong, Nikhil Rodrigo",
        "Semester: Fall 2026",
        "Group Number: 11",
        "Client Group: 8",
        "Coordinator: Adam Dalwai",
        "Name of the Guide: Dr. Tushara Sadasivuni",
    ]:
        line = DOC.add_paragraph(text)
        line.alignment = WD_ALIGN_PARAGRAPH.LEFT
        line.paragraph_format.space_after = Pt(2)
    DOC.add_paragraph().paragraph_format.space_after = Pt(20)
    DOC.add_paragraph("Date: (10/05/2026)")
    heading("TABLE OF CONTENTS", 1, True)
    toc_block()


def toc_block():
    """TOC field whose cached result lists the headings and pages found by the previous pass."""
    entries = json.loads(TOC_FILE.read_text()) if TOC_FILE.exists() else []
    paragraphs = []
    for level, text, page in entries or [[1, "", ""]]:
        paragraph = DOC.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = Inches(0.3 * (level - 1))
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.7), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        run = paragraph.add_run(f"{text}\t{page}" if text else "")
        run.bold = level == 1
        paragraphs.append(paragraph)
    def fld(paragraph, kind, instruction=None, first=False):
        run = OxmlElement("w:r")
        mark = OxmlElement("w:fldChar")
        mark.set(qn("w:fldCharType"), kind)
        run.append(mark)
        if instruction:
            code = OxmlElement("w:instrText")
            code.set(qn("xml:space"), "preserve")
            code.text = instruction
            run.append(code)
        if first:
            paragraph._p.insert(paragraph._p.index(paragraph._p.find(qn("w:r"))) if paragraph._p.find(qn("w:r")) is not None else len(paragraph._p), run)
        else:
            paragraph._p.append(run)
    head = paragraphs[0]
    anchor = head._p.find(qn("w:r"))
    for kind, instruction in [("begin", None), (None, ' TOC \\o "1-2" \\h \\z \\u '), ("separate", None)]:
        run = OxmlElement("w:r")
        if kind:
            mark = OxmlElement("w:fldChar")
            mark.set(qn("w:fldCharType"), kind)
            run.append(mark)
        else:
            code = OxmlElement("w:instrText")
            code.set(qn("xml:space"), "preserve")
            code.text = instruction
            run.append(code)
        if anchor is not None:
            anchor.addprevious(run)
        else:
            head._p.append(run)
    fld(paragraphs[-1], "end")


def baseline_biographies():
    original = Document(ROOT / "docs/Dalwai_Group11_Sprint1.docx")
    paragraphs = original.paragraphs
    start = next(index for index, paragraph in enumerate(paragraphs) if paragraph.text.strip() == "Software Engineers’ information")
    end = next(index for index in range(start + 1, len(paragraphs)) if paragraphs[index].text.strip() == "Planning and Scheduling")
    for paragraph in paragraphs[start + 1:end]:
        if paragraph.text.strip():
            body(paragraph.text)


def assemble():
    configure()
    cover()
    heading("1.0 INTRODUCTION", 1, True)
    heading("1.1 Software Engineers Information", 2)
    baseline_biographies()
    orientation(True)
    heading("1.2 Planning and Scheduling", 2)
    table(["Assignee Name", "Email address", "Task", "Duration (hours)", "Dependency", "Due Date", "Evaluation"],
          PLAN, [1.5, 2.3, 2.3, 0.8, 1.1, 0.7, 0.8], 9)
    orientation(False)
    heading("1.3 Teamwork Basics", 2)
    bullets([
        "We use Discord and iMessage to communicate and share progress.",
        "Shadmaan Akhand was the Sprint 1 coordinator. Adam Dalwai is the Sprint 2 coordinator and combines the sections and keeps the Sprint 2 checklist.",
        "Each member reviews another member's work before we submit.",
        "The GitHub repository holds the source code and the project board tracks each task as To Do, In Progress, or Done.",
        "We only mark a task Done when the work exists and has been checked.",
    ])
    heading("1.4 Problem Statement", 2)
    body("Tutor Plus is a web application that connects adult students with tutors for virtual and in-person sessions. Students often find tutors, book times, pay, and track their learning in separate tools, so it is hard to see how sessions, practice work, and feedback fit together. Tutor Plus keeps these steps in one place. Students find tutors, book a session, pay for it, prepare for it, and review their progress afterwards. Tutors manage their profile, availability, and students. Administrators manage accounts and keep the platform safe.")
    labeled("New in Sprint 2 (bold and highlighted):")
    body("Our system is a Next.js and TypeScript web application that uses Supabase PostgreSQL and Supabase Auth. It is for adults aged 18 and older. Sessions use a meeting link or an in-person place. Direct chat, peer swaps, native video, and guardian accounts are not part of the project.")
    for text in [
        "In Sprint 2 we wrote ten detailed use cases and requirements, built the database for four of them, and implemented User Management, profile editing, and a session reservation flow. Online payment, tutor search, and the learning tools are designed but not built yet.",
        "A booking is saved as payment pending, and a payment counts only after the payment provider confirms it. Private profile and learning data is never shown publicly.",
    ]:
        run = body(text).runs[0]
        run.bold = True
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    labeled("Use cases specified in this sprint:")
    numbered([name for name, *_ in USE_CASES])
    labeled("Implemented in this sprint:", "UC-01 User Management, UC-02 profile editing, and UC-04 session reservation.")
    labeled("Planned for the next sprint:")
    numbered(BACKLOG)
    heading("1.5 System Requirements", 2)
    body("The system serves adults aged 18 or older. A person can be a Student, a Tutor, or both, and an Administrator role is given only by another Administrator. Private profile and learning data is never shown publicly. Bookings must never overlap, and a payment counts only after the payment provider confirms it.")
    heading("1.5.1 Context Diagram", 2)
    body("Tutor Plus is one system connected to the outside systems it depends on: Supabase for accounts and data, Stripe for payments, an email service for verification and reminders, a location service for approximate distance, and the OpenAI API for the learning assistant planned for a later sprint. Students, Tutors, and Administrators use the system through the web application.")
    figure("context-systems.png", "Context Diagram", 4.4)
    body("Inside Tutor Plus, the ten Sprint 2 use cases are grouped into six subsystems: User and Profile Management (UC-01, UC-02), Tutor Search and Discovery (UC-03), Tutor Scheduling and Booking (UC-04, UC-09), Payment and Refund (UC-06), Learning and Feedback (UC-05, UC-07, UC-08), and Notification (UC-10).")
    figure("context-subsystems-sprint2.png", "Sprint 2 Subsystem Overview", 4.4)
    heading("1.5.2 Activity Diagram", 2, True)
    body("The activity diagram follows a Student booking a tutor, from search to review. The lanes are the Student, Tutor Plus, Stripe in test mode, and the Tutor. The slot is rechecked and held atomically, the booking is saved as payment pending, and it is confirmed only after Tutor Plus verifies the signed payment event. A declined payment or an unavailable slot sends the Student back to retry. After the session, the Tutor records a recap and practice task, and the Student leaves a verified review.")
    figure("activity-student-books-tutor-sprint2.png", "Activity Diagram: Student Books a Tutor", 7.3)
    heading("1.6 Client Group Requirements", 2, True)
    body("Group 8 is our client group. The three requirements below are proposed from our design and are waiting for Group 8 to confirm them.")
    for number, text in CLIENT_REQUIREMENTS:
        labeled(f"{number}:", text)
    table(["Requirement", "Group 8 representative", "Date confirmed", "How confirmed", "Status"],
          [[number, "", "", "", "Waiting for confirmation"] for number, _ in CLIENT_REQUIREMENTS],
          [1.0, 1.7, 1.2, 1.4, 1.3], 9)
    heading("1.7 Changes Since Sprint 1", 2)
    table(["Sprint 1 item", "Change in Sprint 2"], [list(row) for row in SPRINT1_CHANGES], [1.7, 4.9], 10)

    heading("2.0 REQUIREMENTS", 1, True)
    heading("2.1 Use Cases", 2)
    labeled("List of Use Case Name:")
    numbered([name for name, *_ in USE_CASES])
    DOC.add_paragraph()
    for number, (name, actors, description, steps, alternate, exception, pre, post) in enumerate(USE_CASES, 1):
        labeled("Use Case no.:", str(number), True)
        labeled("Use Case Name:", name, True)
        labeled("Actors:", actors, True)
        labeled("Description:", description, True)
        labeled("Main Flow:", keep=True)
        numbered(steps, True)
        labeled("Alternate Path:", alternate, True)
        labeled("Exception Path:", exception, True)
        labeled("Pre-condition:", pre, True)
        labeled("Post-condition:", post)
        DOC.add_paragraph()
    heading("2.2 Requirements", 2, True)
    for number, (name, introduction, inputs, description, outputs) in enumerate(REQUIREMENTS, 1):
        labeled("Requirement number:", str(number))
        labeled("Use Case number:", f"{number} ({USE_CASES[number - 1][0]})")
        labeled("Introduction:", introduction)
        labeled("Inputs:")
        numbered(inputs)
        labeled("Requirements Description:", description)
        labeled("Outputs:")
        numbered(outputs)
        DOC.add_paragraph()
    heading("2.3 Use Case Diagrams", 2, True)
    for number, (case, filename, name) in enumerate([
        (1, "usecase-uc01-sprint2.png", "User Management"),
        (3, "usecase-uc03-sprint2.png", "Tutor Search and Discovery"),
        (4, "usecase-uc04-sprint2.png", "Availability Management and Session Booking"),
        (6, "usecase-uc06-sprint2.png", "Payment and Refund Processing"),
    ], 1):
        labeled("Use Case no.:", str(case))
        labeled("Use Case Diagram no.:", str(number))
        labeled("Use Case Name:", name)
        figure(filename, None, 7.2 if number == 1 else 7.2)
        if number < 4:
            DOC.add_paragraph().paragraph_format.page_break_before = True

    heading("3.0 DATABASE TABLES", 1, True)
    body("The database is PostgreSQL 15 managed through Supabase. The 18 tables are in four groups, one for each of the use cases UC-01, UC-02, UC-04, and UC-06. Supabase manages the login accounts in auth.users, so we do not create that table. The same tables are in the migration files and in sql/sprint2_tables.sql.")
    heading("3.1 Database Dictionary and Relationships", 2)
    attributes = database_rows()
    for index, (title, filename) in enumerate([
        ("UC-01 User Management", "erd-uc01.png"),
        ("UC-02 Manage Student and Tutor Profiles", "erd-uc02.png"),
        ("UC-04 Availability and Booking", "erd-uc04.png"),
        ("UC-06 Payments and Refunds", "erd-uc06.png"),
    ]):
        group = DATABASE_GROUPS[index]
        heading(title, 3)
        figure(filename, f"ER diagram for {title}", 4.3)
        for position, (table_name, purpose) in enumerate(group[1], 1):
            labeled(f"{position}) {table_name}", keep=True)
            labeled("Purpose:", purpose + (" It has no relationships and is not drawn in the ER diagram." if table_name in ("admin_bootstrap_guard", "booking_policy") else ""), keep=True)
            labeled("Attributes:", ", ".join(f"{name} {kind}" + (f" ({key})" if key else "") for name, kind, key in attributes[table_name]))
    heading("3.2 SQL and Database Communication", 2)
    body("A form sends its data to a Next.js server action. The server checks who the user is and what they may do, validates the data, and saves it through a Supabase client that runs as that user. Row level security and limited table permissions protect the tables even if someone calls the database API directly. The database itself rejects overlapping bookings and unauthorized role changes.")
    heading("3.3 Supabase Schema", 2)
    body("These screenshots come from the hosted Supabase project, after the same eight migration files were applied. The public schema has 16 tables. The other two tables, admin_bootstrap_guard and booking_policy, are in the private schema and are not exposed.")
    figure("supabase-tables-list.png", "Supabase table list for the public schema", 6.0, CLOUD)
    figure("supabase-schema-part1.png", "Supabase Schema Visualizer, part 1: roles, audit log, subjects, availability, and tutor profiles", 7.0, CLOUD)
    figure("supabase-schema-part2.png", "Supabase Schema Visualizer, part 2: bookings, payments, refunds, and profiles", 7.6, CLOUD)
    figure("supabase-table-profiles.png", "Columns of the profiles table in Supabase", 6.0, CLOUD)
    figure("supabase-table-user-roles.png", "Columns of the user_roles table in Supabase", 3.0, CLOUD)
    figure("supabase-table-bookings.png", "Columns of the bookings table in Supabase", 7.0, CLOUD)

    orientation(True)
    heading("4.0 CLASS DIAGRAM", 1)
    heading("4.1 UC-01 User Management Classes", 2)
    figure("class-uc01.png", "Class diagram for User Management", 6.3)
    heading("4.2 UC-04 Booking Classes", 2, True)
    figure("class-uc04.png", "Class diagram for Availability and Booking", 6.4)
    orientation(False)

    heading("5.0 BEHAVIORAL MODELING", 1, True)
    body("Sequence diagrams are part of a later sprint.")
    heading("6.0 IMPLEMENTATION", 1)
    heading("6.1 User Management Implementation", 2)
    body("Registration and login are web forms that submit to Next.js server actions. The server checks the email, password, display name, time zone, the 18+ box, and the Student or Tutor choice. Supabase Auth creates the account and a database trigger creates the profile and the first role. There is no Administrator option on the form.")
    body("The user must open the verification email before they can sign in. After sign in, the server reads the account status and roles from the database on every protected page. A session lasts 24 hours and is not extended by refreshing. Sign out ends the session. Protected pages redirect a signed-out visitor to the login page, and the admin page is refused for anyone who is not an Administrator. An Administrator can grant or remove roles and change account status, and each change is recorded in the audit log.")
    heading("6.2 Profile and Reservation", 2)
    body("A verified user opens their profile, changes their display name and time zone, and saves. The server uses the signed-in user's id, so a user can only change their own row. The saved values are read from the database and shown again after a page reload.")
    body("A Student can also reserve a session with a published Tutor. The database checks the tutor's availability, lead time, and price, and saves the booking as payment pending. Two bookings for the same tutor time are blocked by the database, and cancelling a booking frees the time. No payment is taken in this sprint.")
    heading("6.3 Build and Run", 2)
    body("The project builds and runs from a terminal, with no IDE. Node.js 22 or newer, npm, Docker, the Supabase CLI, and Git are needed. The steps are:")
    numbered([
        "Clone the repository: git clone https://github.com/addiee13/tutor-plus-sprint2 and open the folder.",
        "Install the packages: npm ci.",
        "Start local Supabase: supabase start. This runs the database and the local test mailbox.",
        "Apply the database migrations: supabase migration up --local.",
        "Copy .env.example to .env.local and set NEXT_PUBLIC_SUPABASE_URL, NEXT_PUBLIC_SUPABASE_PUBLISHABLE_KEY, and NEXT_PUBLIC_SITE_URL=http://127.0.0.1:3000. The URL and key are printed by supabase start. The .env.local file is not committed.",
        "Build the application: bash scripts/build.sh --fresh. It installs, type checks, lints, and builds.",
        "Start it: npm start -- --hostname 127.0.0.1 --port 3000, then open http://127.0.0.1:3000.",
    ], left=True)
    labeled("Verification email:", "Local Supabase does not send real email. Open the test mailbox at http://127.0.0.1:54324, open the message for the new account, and click the verification link.", left=True)
    labeled("Demo data:", "Register a Student and a Tutor through the sign up page and verify both. To make a published Tutor for the reservation demo, run the seed script in the repository's scripts folder with the Tutor's credentials.", left=True)
    labeled("First Administrator:", "An Administrator cannot register on the sign up page. Set SUPABASE_SERVICE_ROLE_KEY and BOOTSTRAP_ADMIN_USER_ID in .env.local to a verified user, then run node scripts/bootstrap-admin.mjs --confirm once. After that the Administrator grants roles from the admin page.", left=True)
    heading("6.4 Build and Run Check", 2)
    body("We ran the steps above and checked the results.")
    table(["Check", "Result"], [
        ["Fresh install, type check, lint, and production build", "Passed"],
        ["Registration, verification email, login, and logout in the browser", "Passed"],
        ["Profile edit, save, and reload", "Passed"],
        ["Non-administrator refused the admin page", "Passed"],
        ["Reservation saved, shown after reload, and cancelled", "Passed"],
        ["Session expires 24 hours after sign in", "Checked by a database test of the session limit. The browser session expiry test was skipped."],
        ["Hosted deployment loads, and /dashboard sends a signed-out visitor to /login", "Passed"],
    ], [3.6, 3.0], 10)
    body("The application is also deployed. The app runs on Vercel at https://tutorplus-one.vercel.app, and the database is a hosted Supabase project with the same eight migrations. The checks above for sign in, profile, and reservation were run locally. The signed-in flow on the hosted site has not been checked yet.")
    figure("desktop-01-landing.png", "Hosted landing page on Vercel", 4.6, CLOUD)
    figure("desktop-02-signup.png", "Hosted registration page", 4.6, CLOUD)
    figure("desktop-05-protected-redirect.png", "Hosted site: opening /dashboard while signed out shows the sign-in page", 4.6, CLOUD)

    heading("7.0 TESTING", 1)
    body("Test cases are written in a later sprint.")
    heading("8.0 ARCHITECTURAL MODELING", 1)
    body("Architectural views are part of a later sprint.")
    heading("9.0 GITHUB", 1, True)
    labeled("Repository:", "https://github.com/addiee13/tutor-plus-sprint2")
    labeled("Project board:", "https://github.com/users/addiee13/projects/1")
    body("The repository has the Next.js application, the database migrations, sql/sprint2_tables.sql, the build script, and a GitHub Actions workflow that installs, type checks, lints, and builds the project. The workflow passed on both commits, f21c128 and 371d948. The board has the columns To Do, In Progress, and Done.")
    heading("GitHub Screenshots", 2)
    for shot in json.loads((ROOT / "docs/evidence/screenshots.json").read_text()):
        if shot["path"].split("/")[-1].startswith("phone-"):
            continue
        name = Path(shot["path"]).name
        if name.split("-")[1:2] and name.split("-")[1] in SKIP_SHOTS:
            continue
        description = re.sub(r"^Desktop:\s*", "", shot["description"])
        figure(name, description, 6.6, ROOT / "docs/evidence/screenshots", 6.4)
    heading("REFERENCES", 1, True)
    for text in [
        "Tutor Plus Software Design Document, version 1.1.",
        "Group 11 Sprint 1 report, September 2026.",
        "Course lecture slides: requirement engineering, system modeling, and database fundamentals.",
        "Next.js App Router documentation, nextjs.org/docs.",
        "Supabase documentation for Auth, row level security, and local development, supabase.com/docs.",
        "PostgreSQL 15 documentation, postgresql.org/docs/15.",
        "GitHub CLI and GitHub Projects documentation, docs.github.com.",
    ]:
        DOC.add_paragraph(text, style="List Number")
    destination = OUT / "Adam_Dalwai_Group11_Sprint2.docx"
    DOC.save(destination)
    (OUT / "headings.json").write_text(json.dumps(HEADINGS))
    print(json.dumps({"docx": str(destination), "paragraphs": len(DOC.paragraphs), "tables": len(DOC.tables)}))


if __name__ == "__main__":
    assemble()
